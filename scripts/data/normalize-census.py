#!/usr/bin/env python3
"""令和7年国勢調査の人口・世帯・年齢構成を公開用の中間JSONへ正規化する。

総数（人口・世帯数と5年間の増減）は原数値の第1-1表、年齢構成の割合は
「不詳」を補完した参考表の第2-7表を使う。割合を原数値で計算すると、
年齢不詳の扱いで割合が市町間・年次間でぶれるためである。
2020年の割合も同じ方法で補完した遡及集計を使い、5年間で比べる。

各表は市区町村ごとに現在の境域と2000年境域の行を並べているため、
「旧：」ではない行だけを採用する。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from census_sources import COMPARISON_REFERENCE_DATE, RAW_ROOT, REFERENCE_DATE, TITLE

ROOT = Path(__file__).resolve().parents[2]
NATIONAL_CODE = "00000"
PREFECTURE_CODE = "34000"

# 第1-1表の列（0始まり）。見出しを照合してから読む。
POPULATION_COLUMNS = {
    "code": 5,
    "name": 6,
    "population": 7,
    "population_2020": 10,
    "population_change": 11,
    "population_change_rate_percent": 12,
    "households": 16,
    "households_2020": 19,
    "households_change": 20,
    "households_change_rate_percent": 21,
}
POPULATION_HEADERS = {
    7: "人口",
    10: "2020年（令和2年）の人口（組替）",
    11: "5年間の人口増減数",
    12: "5年間の人口増減率",
    16: "世帯数",
    19: "2020年（令和2年）の世帯数（組替）",
    20: "5年間の世帯増減数",
    21: "5年間の世帯増減率",
}
POPULATION_ITEMS = {7: "0_総数", 16: "0_総数"}

# 第2-7表の列（0始まり）。
AGE_COLUMNS = {
    "nationality": 0,
    "sex": 1,
    "code": 7,
    "name": 8,
    "total": 9,
    "unknown": 31,
    "age_0_14": 32,
    "age_15_64": 33,
    "age_65_plus": 34,
    "age_75_plus": 35,
    "average_age": 46,
}
AGE_ITEMS = {
    9: "00_総数",
    31: "22_年齢「不詳」",
    32: "R1_（再掲）15歳未満",
    33: "R2_（再掲）15～64歳",
    34: "R3_（再掲）65歳以上",
    35: "R4_（再掲）75歳以上",
}
AGE_HEADERS = {46: "平均年齢"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-root", type=Path, default=RAW_ROOT)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "data/processed/census/2025/census.json",
    )
    parser.add_argument(
        "--municipalities",
        type=Path,
        default=ROOT / "config/municipalities/hiroshima.json",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_manifest(raw_root: Path) -> dict:
    path = raw_root / "source.json"
    if not path.exists():
        raise FileNotFoundError(f"取得manifestがありません。先にacquire:censusを実行してください: {path}")
    manifest = json.loads(path.read_text(encoding="utf-8"))
    for item in manifest["sources"]:
        file_path = raw_root / item["file"]
        if not file_path.is_file() or sha256(file_path) != item["sha256"]:
            raise ValueError(f"原本のハッシュがmanifestと一致しません: {file_path}")
    return manifest


def read_rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path, read_only=True)
    try:
        return list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()


def find_row(rows: list[tuple], label: str, path: Path) -> tuple:
    for row in rows[:30]:
        if label in row:
            return row
    raise ValueError(f"見出し行「{label}」が見つかりません: {path}")


def check_headers(row: tuple, expected: dict[int, str], path: Path) -> None:
    for index, label in expected.items():
        if index >= len(row) or row[index] != label:
            actual = row[index] if index < len(row) else None
            raise ValueError(
                f"列{index}の見出しが想定と異なります（{label!r} ≠ {actual!r}）: {path}"
            )


def split_label(value: object) -> str:
    text = str(value or "")
    _, separator, label = text.partition("_")
    return label if separator else text


def parse_count(value: object, context: str) -> int:
    if value in (None, "-"):
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value != int(value):
        raise ValueError(f"{context}の値を整数として読めません: {value!r}")
    return int(value)


def parse_number(value: object, context: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{context}の値を数値として読めません: {value!r}")
    return float(value)


def is_current_area(name: str) -> bool:
    return "（旧：" not in name


def parse_population(path: Path, codes: set[str]) -> dict[str, dict]:
    rows = read_rows(path)
    check_headers(find_row(rows, "表章項目", path), POPULATION_HEADERS, path)
    check_headers(find_row(rows, "項目名", path), POPULATION_ITEMS, path)
    columns = POPULATION_COLUMNS
    records: dict[str, dict] = {}
    for row in rows:
        if len(row) <= columns["households_change_rate_percent"]:
            continue
        code = str(row[columns["code"]] or "")
        name = split_label(row[columns["name"]])
        if code not in codes or not is_current_area(name):
            continue
        if code in records:
            raise ValueError(f"第1-1表で地域コード{code}の行が重複しています")
        context = f"第1-1表 {code} {name}"
        record = {
            "name_ja": name,
            "population_total": parse_count(row[columns["population"]], context),
            "population_total_2020": parse_count(row[columns["population_2020"]], context),
            "population_change_5y": parse_count(row[columns["population_change"]], context),
            "households_total": parse_count(row[columns["households"]], context),
            "households_total_2020": parse_count(row[columns["households_2020"]], context),
            "households_change_5y": parse_count(row[columns["households_change"]], context),
        }
        # 増減数と増減率は原表の公表値と再計算値を突き合わせて検算する。
        for prefix in ("population", "households"):
            current = record[f"{prefix}_total"]
            previous = record[f"{prefix}_total_2020"]
            change = record[f"{prefix}_change_5y"]
            if current - previous != change:
                raise ValueError(f"{context}の{prefix}増減数が原表と一致しません")
            rate = change / previous
            published = parse_number(row[columns[f"{prefix}_change_rate_percent"]], context) / 100
            if abs(rate - published) > 1e-6:
                raise ValueError(
                    f"{context}の{prefix}増減率が原表と一致しません: {rate} ≠ {published}"
                )
            record[f"{prefix}_change_rate_5y"] = rate
        records[code] = record
    return records


def parse_age(path: Path, codes: set[str], label: str) -> dict[str, dict]:
    rows = read_rows(path)
    check_headers(find_row(rows, "項目名", path), AGE_ITEMS, path)
    check_headers(find_row(rows, "表章項目", path), AGE_HEADERS, path)
    columns = AGE_COLUMNS
    records: dict[str, dict] = {}
    for row in rows:
        if len(row) <= columns["average_age"]:
            continue
        if row[columns["nationality"]] != "0_国籍総数" or row[columns["sex"]] != "0_総数":
            continue
        code = str(row[columns["code"]] or "")
        name = split_label(row[columns["name"]])
        if code not in codes or not is_current_area(name):
            continue
        if code in records:
            raise ValueError(f"{label}で地域コード{code}の行が重複しています")
        context = f"{label} {code} {name}"
        total = parse_count(row[columns["total"]], context)
        unknown = parse_count(row[columns["unknown"]], context)
        groups = {
            key: parse_count(row[columns[key]], context)
            for key in ("age_0_14", "age_15_64", "age_65_plus", "age_75_plus")
        }
        # 不詳補完値でも、補完できなかった少数の年齢不詳が残る地域がある。
        # 割合の分母は年齢が分かる人口とし、残った不詳は人数を別に持つ。
        known = total - unknown
        if groups["age_0_14"] + groups["age_15_64"] + groups["age_65_plus"] != known:
            raise ValueError(f"{context}の年齢3区分の合計が年齢判明人口と一致しません")
        if known <= 0:
            raise ValueError(f"{context}の年齢判明人口が0以下です")
        records[code] = {
            "population_total": total,
            "age_unknown": unknown,
            "shares": {key: value / known for key, value in groups.items()},
            "average_age": parse_number(row[columns["average_age"]], context),
        }
    return records


def area_level(code: str) -> str:
    if code == NATIONAL_CODE:
        return "national"
    if code == PREFECTURE_CODE:
        return "prefecture"
    return "municipality"


def main() -> int:
    args = parse_args()
    raw_root = args.raw_root.resolve()
    manifest = read_manifest(raw_root)
    sources = {item["id"]: item for item in manifest["sources"]}

    municipalities = json.loads(args.municipalities.read_text(encoding="utf-8"))
    municipality_codes = [str(item["code"]) for item in municipalities]
    codes = [NATIONAL_CODE, PREFECTURE_CODE, *municipality_codes]
    code_set = set(codes)

    population = parse_population(raw_root / sources["population_households"]["file"], code_set)
    age_2025 = parse_age(raw_root / sources["age_2025"]["file"], code_set, "2025年第2-7表")
    age_2020 = parse_age(raw_root / sources["age_2020"]["file"], code_set, "2020年第2-7表")

    entries = []
    for code in codes:
        missing = [
            label
            for label, records in (("第1-1表", population), ("2025年第2-7表", age_2025), ("2020年第2-7表", age_2020))
            if code not in records
        ]
        if missing:
            raise ValueError(f"地域コード{code}の行がありません: {', '.join(missing)}")
        totals = population[code]
        # 不詳補完値でも総人口は原数値と同じ。違えば別の地域や境域を読んでいる。
        if age_2025[code]["population_total"] != totals["population_total"]:
            raise ValueError(f"地域コード{code}の2025年人口が第1-1表と第2-7表で一致しません")
        if age_2020[code]["population_total"] != totals["population_total_2020"]:
            raise ValueError(f"地域コード{code}の2020年人口が第1-1表と遡及集計で一致しません")
        entries.append(
            {
                "area_code": code,
                "area_level": area_level(code),
                **totals,
                "age_shares": age_2025[code]["shares"],
                "age_shares_2020": age_2020[code]["shares"],
                "age_unknown_population": age_2025[code]["age_unknown"],
                "age_unknown_population_2020": age_2020[code]["age_unknown"],
                "average_age": age_2025[code]["average_age"],
                "average_age_2020": age_2020[code]["average_age"],
            }
        )

    payload = {
        "schema_version": 1,
        "dataset": "census",
        "title": TITLE,
        "reference_date": REFERENCE_DATE,
        "comparison_reference_date": COMPARISON_REFERENCE_DATE,
        "total_value_type": "original",
        "share_value_type": "imputed",
        "acquired_at": manifest["acquired_at"],
        "index_url": manifest["index_url"],
        "sources": [
            {
                "id": item["id"],
                "title": item["title"],
                "table_number": item["table_number"],
                "value_type": item["value_type"],
                "url": item["url"],
                "file": item["file"],
                "sha256": item["sha256"],
            }
            for item in manifest["sources"]
        ],
        "entries": entries,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"国勢調査の正規化OK: {args.output} / {len(entries)}地域 / 基準日{REFERENCE_DATE}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (FileNotFoundError, ValueError, KeyError) as error:
        print(f"国勢調査の正規化失敗: {error}", file=sys.stderr)
        raise SystemExit(1)
